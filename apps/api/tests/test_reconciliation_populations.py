"""TC-334/344: shift balances, calendar activity and real physical counts."""

from datetime import timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from restaurant_os import models
from restaurant_os.operations import BusinessError, _cash_summary_for_shift
from restaurant_os.reconciliation_reports import (
    export_reconciliation_workbook,
    get_branch_daily_reconciliation,
    get_multi_branch_consolidated_report,
)
from test_cash_ledger import BRANCH_A, CASHIER_ID, NOW, ORG_ID, SHIFT_ID, _new_session
from test_system_remediation_reports import cash_row, prepare, purchase_row


@pytest.fixture
def session():
    engine, session = _new_session()
    prepare(session)
    yield session
    session.close()
    engine.dispose()


def report(session, day=0):
    return get_branch_daily_reconciliation(
        session, BRANCH_A, (NOW + timedelta(days=day)).date().isoformat()
    )


def close(session, shift_id=SHIFT_ID, *, counted=None, when=None):
    shift = dict(
        session.execute(models.cash_shifts.select().where(models.cash_shifts.c.id == shift_id))
        .mappings()
        .one()
    )
    summary = _cash_summary_for_shift(session, shift)
    when = when or NOW + timedelta(days=2)
    session.execute(
        models.cash_shifts.update()
        .where(models.cash_shifts.c.id == shift_id)
        .values(status="OPERATIVELY_CLOSED" if counted is None else "CLOSED", closed_at=when)
    )
    if counted is None:
        session.execute(
            models.cash_shift_closures.insert().values(
                id="closure-" + shift_id,
                organization_id=ORG_ID,
                branch_id=BRANCH_A,
                cash_shift_id=shift_id,
                register_code_snapshot=shift["register_code"],
                closed_by_user_id=CASHIER_ID,
                summary_snapshot=summary,
                closed_at=when,
                created_at=when,
            )
        )
    else:
        session.execute(
            models.cash_shift_cuts.insert().values(
                id="cut-" + shift_id,
                organization_id=ORG_ID,
                branch_id=BRANCH_A,
                cash_shift_id=shift_id,
                sales_total_cents=summary["sales_total_cents"],
                payment_total_cents=summary["payment_total_cents"],
                cash_payment_total_cents=summary["cash_payment_cents"],
                opening_cash_cents=summary["opening_cash_cents"],
                expected_cash_cents=summary["expected_cash_cents"],
                counted_cash_cents=counted,
                difference_cents=counted - summary["expected_cash_cents"],
                status="FINAL",
                created_at=when,
            )
        )
    session.commit()
    return summary


def test_cross_midnight_balance_uses_whole_shift_activity_uses_event_day(session):
    purchase_row(session, confirmed_at=NOW + timedelta(days=1))
    cash_row(
        session, "aud-withdrawal", "withdrawal", "PURCHASE", occurred_at=NOW + timedelta(days=1)
    )
    session.commit()
    first, second = report(session), report(session, 1)
    assert first["balance"]["expected_cash_in_register"] == Decimal("1700")
    assert second["balance"]["expected_cash_in_register"] == 0
    assert first["activity"]["totals"]["supplier_expenses"] == 0
    assert second["activity"]["totals"]["supplier_expenses"] == 300
    consolidated = get_multi_branch_consolidated_report(
        session, NOW.date().isoformat(), (NOW + timedelta(days=1)).date().isoformat(), BRANCH_A
    )
    assert consolidated["summary"]["total_expected_cash"] == 1700
    assert consolidated["activity"]["totals"]["supplier_expenses"] == 300


@pytest.mark.parametrize("closed", [False, True])
def test_opening_and_operational_close_are_not_physical_counts(session, closed):
    if closed:
        close(session)
    result = report(session)
    assert result["balance"]["physical_cash_count"] is None
    assert result["balance"]["difference"] is None
    assert result["physical_count"]["status"] == "PENDING"
    assert result["physical_count"]["pending_shift_ids"] == [SHIFT_ID]


@pytest.mark.parametrize("counted", [0, 200000, 210000])
def test_explicit_legacy_whole_shift_count_preserves_zero(session, counted):
    close(session, counted=counted)
    result = report(session)
    assert result["balance"]["physical_cash_count"] == Decimal(counted) / 100
    assert result["balance"]["difference"] == Decimal(counted - 200000) / 100
    assert result["physical_count"]["status"] == "COUNTED"


def test_multiple_shifts_same_register_and_one_pending_do_not_certify_total(session):
    close(session, counted=0, when=NOW + timedelta(hours=1))
    shift = dict(
        session.execute(models.cash_shifts.select().where(models.cash_shifts.c.id == SHIFT_ID))
        .mappings()
        .one()
    )
    session.execute(
        models.cash_shifts.insert().values(
            **{
                **shift,
                "id": "second-shift",
                "status": "OPEN",
                "closed_at": None,
                "opened_at": NOW + timedelta(hours=2),
                "opening_cash_cents": 10000,
            }
        )
    )
    session.commit()
    result = report(session)
    assert result["balance"]["expected_cash_in_register"] == 2100
    assert result["balance"]["physical_cash_count"] is None
    assert result["balance"]["difference"] is None
    assert set(result["population"]["shift_ids"]) == {SHIFT_ID, "second-shift"}
    assert result["physical_count"]["counted_shift_ids"] == [SHIFT_ID]
    consolidated = get_multi_branch_consolidated_report(
        session, NOW.date().isoformat(), NOW.date().isoformat(), BRANCH_A
    )
    assert consolidated["summary"]["physical_cash_count"] is None
    assert consolidated["summary"]["difference"] is None


def test_calendar_refund_preserves_positive_activity_and_frozen_snapshot(session):
    purchase_row(session, cancelled_at=NOW + timedelta(days=1))
    cash_row(session, "aud-withdrawal", "withdrawal", "PURCHASE")
    before = report(session)["activity"]
    cash_row(
        session,
        "return",
        "deposit",
        "PURCHASE_CANCELLATION",
        occurred_at=NOW + timedelta(days=1),
        linked="aud-withdrawal",
    )
    frozen = close(session)
    assert report(session)["activity"] == before
    assert report(session, 1)["activity"]["totals"]["supplier_expenses"] == -300
    assert report(session)["balance"]["expected_cash_in_register"] == 2000
    assert (
        session.execute(models.cash_shift_closures.select()).mappings().one()["summary_snapshot"]
        == frozen
    )


@pytest.mark.parametrize(
    "fault", ["late_effect", "frozen_expected", "count_difference", "cut_scope", "cut_status"]
)
def test_corrupt_closed_population_fails_without_rewriting_snapshot(session, fault):
    frozen = close(session, counted=0) if fault.startswith(("count", "cut")) else close(session)
    if fault == "late_effect":
        cash_row(
            session, "late", "deposit", "MANUAL", amount=100, occurred_at=NOW + timedelta(days=3)
        )
    elif fault == "frozen_expected":
        session.execute(
            models.cash_shift_closures.update().values(
                summary_snapshot={**frozen, "expected_cash_cents": 999}
            )
        )
    elif fault == "count_difference":
        session.execute(models.cash_shift_cuts.update().values(difference_cents=1))
    elif fault == "cut_scope":
        session.execute(models.cash_shift_cuts.update().values(branch_id="foreign-branch"))
    else:
        session.execute(models.cash_shift_cuts.update().values(status="VOIDED"))
    session.commit()
    before = [dict(row) for row in session.execute(models.cash_shift_closures.select()).mappings()]
    with pytest.raises(BusinessError) as error:
        report(session)
    assert error.value.code == "reconciliation_integrity_conflict"
    assert [
        dict(row) for row in session.execute(models.cash_shift_closures.select()).mappings()
    ] == before


@pytest.mark.parametrize("counted", [None, 0])
def test_excel_distinguishes_missing_count_from_real_zero(session, counted):
    close(session, counted=counted)
    book = load_workbook(export_reconciliation_workbook(session, BRANCH_A, NOW.month, NOW.year))
    assert book["Resumen"]["B12"].value == counted
    assert book["Resumen"]["B13"].value == (None if counted is None else -2000)
    assert "apertura" in book["Resumen"]["A2"].value.lower()
    assert "Actividad calendario" in book.sheetnames


def test_no_shifts_is_empty_not_a_count_of_zero(session):
    result = report(session, 10)
    assert result["physical_count"]["status"] == "EMPTY"
    assert result["balance"]["physical_cash_count"] is None
    assert result["balance"]["difference"] is None


@pytest.mark.parametrize(
    "fault", ["foreign_shift", "late_calendar_effect", "foreign_concept", "foreign_order"]
)
def test_activity_rejects_foreign_or_post_close_links_without_leaking_names(session, fault):
    from test_cash_ledger import BRANCH_B

    if fault == "foreign_shift":
        session.execute(
            models.cash_shifts.update()
            .where(models.cash_shifts.c.id == SHIFT_ID)
            .values(branch_id=BRANCH_B)
        )
        cash_row(session, "activity-fault", "deposit", "MANUAL", amount=100)
        day = 0
    elif fault == "late_calendar_effect":
        close(session)
        cash_row(
            session,
            "activity-fault",
            "deposit",
            "MANUAL",
            amount=100,
            occurred_at=NOW + timedelta(days=3),
        )
        day = 3
    elif fault == "foreign_order":
        session.execute(
            models.orders.insert().values(
                id="foreign-order",
                organization_id="foreign-organization",
                branch_id=BRANCH_A,
                cash_shift_id=SHIFT_ID,
                folio="PRIVATE",
                channel="POS",
                status="CLOSED",
                total_cents=100,
                customer_snapshot={"name": "PRIVATE_CUSTOMER"},
                created_at=NOW,
            )
        )
        session.execute(
            models.payments.insert().values(
                id="foreign-order-payment",
                organization_id=ORG_ID,
                branch_id=BRANCH_A,
                cash_shift_id=SHIFT_ID,
                order_id="foreign-order",
                method="cash",
                status="CONFIRMED",
                amount_cents=100,
                confirmed_at=NOW,
                created_at=NOW,
            )
        )
        day = 0
    else:
        session.execute(
            models.cash_movement_concepts.insert().values(
                id="foreign-concept",
                organization_id="foreign-organization",
                code="PRIVATE",
                created_by_user_id=CASHIER_ID,
                created_at=NOW,
            )
        )
        session.execute(
            models.cash_movement_concept_versions.insert().values(
                id="foreign-concept-version",
                concept_id="foreign-concept",
                version=1,
                name="PRIVATE_OTHER_ORG_CONCEPT",
                allowed_movement_type="withdrawal",
                requires_reference=True,
                requires_evidence=True,
                valid_from=NOW,
                created_by_user_id=CASHIER_ID,
                created_at=NOW,
            )
        )
        cash_row(session, "activity-fault", "withdrawal", "MANUAL", amount=100)
        session.execute(
            models.cash_movements.update()
            .where(models.cash_movements.c.id == "activity-fault")
            .values(concept_id="foreign-concept", concept_version_id="foreign-concept-version")
        )
        day = 0
    session.commit()
    date = (NOW + timedelta(days=day)).date().isoformat()
    for read in (
        lambda: report(session, day),
        lambda: get_multi_branch_consolidated_report(session, date, date, BRANCH_A),
        lambda: export_reconciliation_workbook(session, BRANCH_A, NOW.month, NOW.year),
    ):
        with pytest.raises(BusinessError) as error:
            read()
        assert error.value.code == "reconciliation_integrity_conflict"
        assert "PRIVATE" not in str(error.value)


def test_user_count_is_not_equivalent_to_whole_shift_population(session):
    session.execute(
        models.user_cash_cuts.insert().values(
            id="partial-user-count",
            organization_id=ORG_ID,
            branch_id=BRANCH_A,
            cash_shift_id=SHIFT_ID,
            register_code_snapshot="CAJA-01",
            cashier_user_id=CASHIER_ID,
            timezone="UTC",
            period_start=NOW,
            period_end=NOW + timedelta(hours=1),
            status="FINALIZED",
            opening_cash_cents=200000,
            cash_payment_cents=0,
            deposit_cents=0,
            withdrawal_cents=0,
            expected_cash_cents=200000,
            counted_cash_cents=200000,
            difference_cents=0,
            created_by_user_id=CASHIER_ID,
            finalized_by_user_id=CASHIER_ID,
            created_at=NOW,
            counted_at=NOW,
            finalized_at=NOW,
        )
    )
    close(session)
    assert report(session)["physical_count"]["status"] == "PENDING"
    assert report(session)["balance"]["difference"] is None


@pytest.mark.parametrize("day,hours", [("2026-03-08", 23), ("2026-11-01", 25)])
def test_activity_half_open_bounds_and_shift_selection_through_dst(session, day, hours):
    from restaurant_os.reconciliation_reports import _branch_day_bounds_utc

    session.execute(
        models.branches.update()
        .where(models.branches.c.id == BRANCH_A)
        .values(timezone="America/New_York")
    )
    start, end = _branch_day_bounds_utc(session, BRANCH_A, day)
    session.execute(
        models.cash_shifts.update()
        .where(models.cash_shifts.c.id == SHIFT_ID)
        .values(opened_at=start)
    )
    cash_row(session, "at-start", "deposit", "MANUAL", amount=100, occurred_at=start)
    cash_row(session, "at-end", "deposit", "MANUAL", amount=200, occurred_at=end)
    session.commit()
    result = get_branch_daily_reconciliation(session, BRANCH_A, day)
    assert end - start == timedelta(hours=hours)
    assert result["population"]["shift_ids"] == [SHIFT_ID]
    assert result["balance"]["cash_deposits"] == 3
    assert result["activity"]["totals"]["cash_deposits"] == 1
    next_day = get_branch_daily_reconciliation(
        session,
        BRANCH_A,
        (end.astimezone(__import__("zoneinfo").ZoneInfo("America/New_York"))).date().isoformat(),
    )
    assert next_day["population"]["shift_ids"] == []
    assert next_day["activity"]["totals"]["cash_deposits"] == 2


@pytest.mark.parametrize("counted", [None, 0])
def test_live_response_matches_nullable_json_contract(session, counted):
    import json
    from copy import deepcopy
    from pathlib import Path

    from jsonschema import Draft202012Validator

    close(session, counted=counted)
    for kind, payload in [
        ("daily", report(session)),
        (
            "consolidated",
            get_multi_branch_consolidated_report(
                session, NOW.date().isoformat(), NOW.date().isoformat(), BRANCH_A
            ),
        ),
    ]:
        path = (
            Path(__file__).resolve().parents[3]
            / "packages/contracts/schemas"
            / f"reconciliation-{kind}-v2.schema.json"
        )
        schema = json.loads(path.read_text())
        validator = Draft202012Validator(schema)
        from restaurant_os.reconciliation_reports import reconciliation_wire

        encoded = reconciliation_wire(payload)
        validator.validate(encoded)
        invalid = deepcopy(encoded)
        key = "balance" if kind == "daily" else "summary"
        invalid[key]["physical_cash_count"] = 0 if counted is None else None
        assert list(validator.iter_errors(invalid))


@pytest.mark.parametrize("date", ["invalid", "2026-02-30", "2026-1-1", "9999-12-31"])
def test_invalid_report_dates_fail_explicitly(session, date):
    with pytest.raises(BusinessError) as error:
        get_branch_daily_reconciliation(session, BRANCH_A, date)
    assert error.value.code == "report_period_invalid"


def test_invalid_branch_timezone_does_not_invent_a_local_day(session):
    session.execute(
        models.branches.update()
        .where(models.branches.c.id == BRANCH_A)
        .values(timezone="Invalid/Timezone")
    )
    with pytest.raises(BusinessError) as error:
        report(session)
    assert error.value.code == "report_timezone_invalid"


def test_empty_corporate_population_keeps_complete_activity_contract(session, monkeypatch):
    import restaurant_os.reconciliation_reports as reconciliation

    monkeypatch.setattr(reconciliation, "ORGANIZATION_ID", "synthetic-empty-organization")
    result = get_multi_branch_consolidated_report(
        session, NOW.date().isoformat(), NOW.date().isoformat()
    )
    assert result["branches"] == []
    assert result["physical_count"]["status"] == "EMPTY"
    assert result["summary"]["physical_cash_count"] is None
    assert result["summary"]["difference"] is None
    assert result["activity"]["totals"] == {
        key: Decimal("0")
        for key in (
            "total_sales_with_tax",
            "card_payments",
            "transfer_payments",
            "credit_sales",
            "cash_sales",
            "supplier_expenses",
            "fixed_expenses",
            "cash_withdrawals",
            "cash_deposits",
        )
    }
